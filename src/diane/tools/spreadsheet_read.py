"""Built-in tool: read Excel (.xlsx) spreadsheet files into structured data."""

from __future__ import annotations

import os

from diane.tools._workspace import validate_path

TOOL_METADATA = {
    "tool_id": "spreadsheet_read",
    "description": (
        "Read an Excel (.xlsx) spreadsheet file into structured rows. "
        "Returns headers and rows as a list of dicts, matching csv_read output. "
        "Requires the openpyxl package."
    ),
    "parameters": {
        "type": "object",
        "properties": {
            "path": {
                "type": "string",
                "description": "Path to the .xlsx file.",
            },
            "sheet": {
                "type": "string",
                "description": "Sheet name. Omit to read the first (active) sheet.",
            },
            "max_rows": {
                "type": "integer",
                "description": "Maximum number of data rows to return.",
                "default": 1000,
            },
        },
        "required": ["path"],
    },
    "examples": [
        {
            "input": {"path": "report.xlsx", "max_rows": 5},
            "output": {
                "headers": ["name", "revenue"],
                "rows": [{"name": "Q1", "revenue": 1000}],
                "row_count": 1,
                "column_count": 2,
                "truncated": False,
            },
        },
    ],
    "category": "data",
    "returns": "dict with headers, rows (list of dicts), row_count, column_count, truncated",
}

MAX_FILE_SIZE = 100_000_000  # 100 MB


async def spreadsheet_read(
    path: str,
    sheet: str | None = None,
    max_rows: int = 1000,
    **_kwargs,
) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError(
            "openpyxl is required for spreadsheet_read. "
            "Install: pip install 'dan[spreadsheet]'"
        )

    resolved = validate_path(path)

    if not os.path.isfile(resolved):
        raise FileNotFoundError(f"File not found: '{path}'")

    ext = os.path.splitext(resolved)[1].lower()
    if ext not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        raise ValueError(f"Unsupported format '{ext}'. Supported: .xlsx, .xlsm, .xltx, .xltm")

    size = os.path.getsize(resolved)
    if size > MAX_FILE_SIZE:
        raise ValueError(f"File '{path}' is {size:,} bytes (limit {MAX_FILE_SIZE:,}).")

    wb = load_workbook(resolved, read_only=True, data_only=True)
    try:
        if sheet:
            if sheet not in wb.sheetnames:
                raise ValueError(
                    f"Sheet '{sheet}' not found. Available: {wb.sheetnames}"
                )
            ws = wb[sheet]
        else:
            ws = wb.active
            if ws is None:
                raise ValueError("Workbook has no active sheet.")

        row_iter = ws.iter_rows(values_only=True)

        header_row = next(row_iter, None)
        if header_row is None:
            return {
                "headers": [],
                "rows": [],
                "row_count": 0,
                "total_rows": 0,
                "column_count": 0,
                "truncated": False,
            }

        headers = [str(c) if c is not None else f"column_{i}" for i, c in enumerate(header_row)]

        rows: list[dict] = []
        total = 0
        for values in row_iter:
            if all(v is None for v in values):
                continue
            total += 1
            if total <= max_rows:
                row = {}
                for i, h in enumerate(headers):
                    val = values[i] if i < len(values) else None
                    row[h] = val
                rows.append(row)

        return {
            "headers": headers,
            "rows": rows,
            "row_count": len(rows),
            "total_rows": total,
            "column_count": len(headers),
            "truncated": total > max_rows,
        }
    finally:
        wb.close()
